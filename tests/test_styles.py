"""Styles and chapter contexts survive editing and format boundaries."""

from dataclasses import asdict, replace
from pathlib import Path
from io import BytesIO
import zipfile
import json

import fitz
import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from bamboo import (
    Block,
    Book,
    Inline,
    PRESETS,
    SectionSpec,
    TextStyle,
    EditorSession,
    compose,
    from_dict,
    render,
    BambooError,
)
from bamboo.styles import context_at, page_style_presets
from bamboo.importers import import_docx


def mixed_book():
    return Book(
        "篇章測試",
        (
            Block((Inline("竪排正文"),), style="preface"),
            Block(
                (Inline("横排正文"),),
                section=SectionSpec(
                    name="横排", profile=PRESETS["horizontal"], page_number_start=7
                ),
            ),
            Block((Inline("讀書養心", "seal", seal_style="white"),), style="author"),
        ),
        profile=PRESETS["blank"],
    )


def test_chapter_geometry_and_source_mapping():
    book = mixed_book()
    layout = compose(book)
    assert len(layout.pages) == 2
    assert [p.folio for p in layout.pages] == [1, 7]
    assert [p.profile.vertical for p in layout.pages] == [True, False]
    assert {g.block for g in layout.pages[1].glyphs if g.block >= 0} == {1, 2}
    s = EditorSession(book)
    for pi, p in enumerate(s.layout().pages):
        g = next(g for g in p.glyphs if g.block >= 0)
        hit = s.hit_test(pi, g.x + g.width * 0.1, g.y + g.height * 0.1)
        assert hit.block_id == s.block_ids[g.block]
        caret = s.caret({"block_id": hit.block_id, "offset": 1})
        assert (caret["y1"] == caret["y2"]) == p.profile.vertical


def test_current_chapter_profile_edit_and_undo():
    s = EditorSession(mixed_book())
    s.select({"block_id": s.block_ids[1], "offset": 0})
    s.dispatch({"type": "set_profile", "values": {"font_size": 14}})
    assert s.book.profile.font_size == PRESETS["blank"].font_size
    assert context_at(s.book, 1).profile.font_size == 14
    s.dispatch({"type": "undo"})
    assert context_at(s.book, 1).profile.font_size == PRESETS["horizontal"].font_size
    s.dispatch({"type": "set_section", "values": {"volume": "二卷", "author": "編者"}})
    assert context_at(s.book, 2).author == "編者"
    s.dispatch({"type": "clear_section"})
    assert context_at(s.book, 1).profile.vertical


def test_reusable_style_updates_every_use():
    s = EditorSession(
        Book(
            "樣式",
            (
                Block((Inline("甲乙"),), style="preface"),
                Block((Inline("丙丁"),), style="preface"),
            ),
            profile=PRESETS["blank"],
        )
    )
    s.dispatch(
        {
            "type": "define_style",
            "values": {"key": "preface", "font_scale": 0.7, "ink": "#b32624"},
        }
    )
    assert {g.size for p in s.layout().pages for g in p.glyphs if g.block >= 0} == {
        s.book.profile.font_size * 0.7
    }
    assert {g.color for p in s.layout().pages for g in p.glyphs if g.block >= 0} == {
        "#b32624"
    }
    s.dispatch({"type": "undo"})
    assert not s.book.styles


def test_poetry_linebreak_and_enter_inherit_style():
    s = EditorSession()
    s.dispatch({"type": "apply_style", "style": "poetry"})
    s.dispatch({"type": "insert_text", "text": "春眠不覺曉"})
    s.dispatch({"type": "insert_linebreak"})
    s.dispatch({"type": "insert_text", "text": "處處聞啼鳥"})
    assert len(s.book.blocks) == 1
    assert s.book.blocks[0].text == "春眠不覺曉\n處處聞啼鳥"
    glyphs = [g for p in s.layout().pages for g in p.glyphs if g.block == 0]
    assert glyphs[0].x > glyphs[5].x and glyphs[0].y == glyphs[5].y
    s.dispatch({"type": "split_paragraph"})
    assert s.book.blocks[1].style == "poetry"


def test_soft_break_updates_annotation_anchor():
    s = EditorSession()
    s.dispatch({"type": "insert_text", "text": "甲乙丙丁"})
    at = {"block_id": s.block_ids[0], "offset": 2}
    s.select(at)
    s.dispatch({"type": "add_annotation", "text": "旁批"})
    s.select({"block_id": s.block_ids[0], "offset": 0})
    s.dispatch({"type": "insert_linebreak"})
    assert s.book.annotations[0].offset == 3
    assert len(s.book.blocks) == 1


def test_cover_optional_and_body_untouched():
    s = EditorSession()
    s.dispatch({"type": "insert_text", "text": "正文"})
    original = s.block_ids[0]
    s.dispatch(
        {
            "type": "insert_cover",
            "title": "简牍讀書集",
            "subtitle": "一卷",
            "border": "none",
        }
    )
    assert s.block_ids[2] == original
    assert len(s.layout().pages) == 2
    assert not s.layout().pages[0].lines
    assert s.book.blocks[2].text == "正文"
    restored = EditorSession.restore(
        json.loads(json.dumps({"editor_schema": 1, **s.state()}))
    )
    assert restored.book == s.book
    s.dispatch({"type": "undo"})
    assert list(s.block_ids) == [original]


def test_overlarge_style_retains_edit_with_diagnostic():
    s = EditorSession()
    s.dispatch({"type": "insert_text", "text": "保留文字"})
    s.dispatch({"type": "define_style", "values": {"key": "body", "font_scale": 3}})
    assert s.layout().pages and s.issues
    assert s.book.blocks[0].text == "保留文字"


@pytest.mark.parametrize("preset", page_style_presets())
def test_page_presets_have_valid_geometry(preset):
    name, p = page_style_presets()[preset]
    assert compose(Book(name, (Block((Inline("天地玄黃宇宙洪荒"),)),), profile=p)).pages


def test_json_roundtrip_styles_and_sections():
    book = replace(mixed_book(), styles=(TextStyle("preface", "序言", ink="#c52c27"),))
    assert from_dict(json.loads(json.dumps(asdict(book)))) == book


def test_three_exports_native_sections_cover_and_seals(tmp_path):
    s = EditorSession(mixed_book())
    s.dispatch({"type": "insert_cover", "title": "简牍讀書集", "subtitle": "一卷"})
    result = render(s.book, tmp_path)
    doc = Document(result.files["docx"])
    assert len(doc.sections) == 3
    assert [
        s._sectPr.find(qn("w:textDirection")).get(qn("w:val")) for s in doc.sections
    ] == ["tbRl", "tbRl", "lrTb"]
    with fitz.open(result.files["pdf"]) as pdf:
        assert len(pdf) == 3
        assert [(round(p.rect.width), round(p.rect.height)) for p in pdf] == [
            (round(p.profile.width), round(p.profile.height)) for p in s.layout().pages
        ]
    html = Path(result.files["html"]).read_text()
    assert "@page leaf1" in html and "@page leaf3" in html
    root = doc._element
    assert (
        len([n for n in root.iter() if n.get("id", "").startswith("BambooSeal")]) == 1
    )
    assert any(n.get("id", "") == "BambooCover0" for n in root.iter())
    imported, warnings = import_docx(Path(result.files["docx"]).read_bytes())
    assert [b.text for b in imported.book.blocks] == [b.text for b in s.book.blocks]
    assert len(imported.layout().pages) == 3
    assert imported.book.blocks[-1].inlines[0].seal_style == "white"
    # A Word edit must win over the embedded source snapshot.
    raw = BytesIO()
    with (
        zipfile.ZipFile(result.files["docx"]) as original,
        zipfile.ZipFile(raw, "w") as changed,
    ):
        for name in original.namelist():
            data = original.read(name)
            if name == "word/document.xml":
                data = data.replace("简牍讀書集".encode(), "新編讀書集".encode())
            changed.writestr(name, data)
    changed, _ = import_docx(raw.getvalue())
    assert changed.book.blocks[0].text == "新編讀書集"


def test_facsimile_retains_each_page_size(tmp_path):
    book = mixed_book()
    result = render(book, tmp_path, formats=("docx",), docx_mode="facsimile", dpi=72)
    assert [
        round(s.page_width.pt) for s in Document(result.files["docx"]).sections
    ] == [round(p.profile.width) for p in compose(book).pages]


NEW_PAGE_STYLES = {
    "single-ink",
    "blue-manuscript",
    "warm-edition",
    "large-print",
    "pocket-book",
    "annotation-page",
    "colophon-page",
    "sutra-page",
    "horizontal-study",
    "horizontal-columns",
}


def test_page_style_catalog_has_ten_new_distinct_layouts():
    from bamboo.styles import PAGE_STYLE_DESCRIPTIONS

    presets = page_style_presets()
    assert NEW_PAGE_STYLES <= presets.keys()
    assert len(presets) >= 16
    assert presets.keys() == PAGE_STYLE_DESCRIPTIONS.keys()
    assert len(
        {json.dumps(asdict(p), sort_keys=True) for _, p in presets.values()}
    ) == len(presets)


@pytest.mark.parametrize("key", page_style_presets())
def test_all_page_styles_export_editable_word_and_fixed_formats(key, tmp_path):
    from bamboo.styles import page_style_sample

    book = page_style_sample(key)
    result = render(book, tmp_path, basename=key)
    p = book.profile
    assert result.pages == 1
    native = Document(result.files["docx"])
    section = native.sections[0]
    assert abs(section.page_width.pt - p.width) < 0.1
    assert abs(section.page_height.pt - p.height) < 0.1
    assert section._sectPr.find(qn("w:textDirection")).get(qn("w:val")) == (
        "tbRl" if p.vertical else "lrTb"
    )
    assert section._sectPr.find(qn("w:cols")).get(qn("w:num")) == str(
        1 if p.vertical else p.panels
    )
    assert all(b.text for b in book.blocks)
    assert len(native.paragraphs) == len(book.blocks)
    with fitz.open(result.files["pdf"]) as pdf:
        assert len(pdf) == 1
        assert abs(pdf[0].rect.width - p.width) < 0.1
        assert "读书" in pdf[0].get_text().replace("\n", "")
    assert "<svg" in Path(result.files["html"]).read_text()
    back, warnings = import_docx(Path(result.files["docx"]).read_bytes())
    assert back.book.profile.vertical == p.vertical
    assert len(back.book.blocks) == len(book.blocks)


def test_page_style_changes_preserve_text_ids_and_are_undoable():
    s = EditorSession()
    s.dispatch({"type": "insert_text", "text": "第一段\n第二段"})
    s.select({"block_id": s.block_ids[1], "offset": 0})
    s.dispatch({"type": "add_annotation", "text": "小批", "extent": 3})
    before = s.book
    ids = s.block_ids
    s.dispatch({"type": "set_section", "preset": "annotation-page", "new": True})
    assert [b.text for b in s.book.blocks] == [b.text for b in before.blocks]
    assert s.book.annotations == before.annotations and s.block_ids == ids
    assert context_at(s.book, 0).profile == before.profile
    assert context_at(s.book, 1).profile == page_style_presets()["annotation-page"][1]
    s.dispatch({"type": "undo"})
    assert s.book == before
