"""Native editable tables for specialist records, with tagged live fields."""

from dataclasses import replace
from docx.shared import Pt
from docx.enum.table import (
    WD_TABLE_ALIGNMENT,
    WD_ROW_HEIGHT_RULE,
    WD_CELL_VERTICAL_ALIGNMENT,
)
from docx.oxml.ns import qn
from .wordflow import element, _run, _section, _styles


def tagged(paragraph, value, font, size, color, tag=None):
    run = _run(paragraph, value, font, size, color)
    if tag:
        sdt = element("w:sdt")
        props = element("w:sdtPr")
        props.append(element("w:tag", val=tag))
        sdt.append(props)
        content = element("w:sdtContent")
        sdt.append(content)
        paragraph._p.remove(run._r)
        content.append(run._r)
        paragraph._p.append(sdt)


def table(doc, widget, font, p, kind):
    t = doc.add_table(rows=widget["rows"], cols=widget["columns"])
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    pr = t._tbl.tblPr
    pr.append(element("w:tblInd", w=0, type="dxa"))
    borders = element("w:tblBorders")
    for name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        borders.append(
            element(
                "w:" + name,
                val="single" if p.border != "none" else "nil",
                sz=5,
                color=p.rule_color[1:],
            )
        )
    pr.append(borders)
    margins = element("w:tblCellMar")
    for side in ["top", "left", "bottom", "right"]:
        margins.append(element("w:" + side, w=80, type="dxa"))
    pr.append(margins)
    for c, w in zip(t.columns, widget["widths"]):
        c.width = Pt(w)
    for ri, row in enumerate(t.rows):
        row.height = Pt(max(8, widget["heights"][ri] - 8))
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row._tr.get_or_add_trPr().append(element("w:cantSplit"))
        for ci, c in enumerate(row.cells):
            c.width = Pt(widget["widths"][ci])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for raw in widget["cells"]:
        if raw.get("rowspan", 1) > 1 or raw.get("colspan", 1) > 1:
            c = t.cell(raw["row"], raw["column"]).merge(
                t.cell(
                    raw["row"] + raw.get("rowspan", 1) - 1,
                    raw["column"] + raw.get("colspan", 1) - 1,
                )
            )
            for extra in list(c._tc.findall(qn("w:p")))[1:]:
                c._tc.remove(extra)
    for raw in widget["cells"]:
        c = t.cell(raw["row"], raw["column"])
        pp = c.paragraphs[0]
        pp.paragraph_format.space_before = pp.paragraph_format.space_after = Pt(0)
        pp.paragraph_format.line_spacing = Pt(raw["size"] * 1.35)
        if raw["vertical"]:
            c._tc.get_or_add_tcPr().append(element("w:textDirection", val="tbRlV"))
        pp._p.get_or_add_pPr().append(element("w:snapToGrid", val="0"))
        tag = (
            f'bamboo:{kind}:{raw["record"]}:{raw["field"]}'
            if raw.get("record") and raw.get("editable", True)
            else None
        )
        tagged(pp, raw["text"], font, raw["size"], raw["color"], tag)
    if not p.vertical:
        t.rows[0]._tr.get_or_add_trPr().append(element("w:tblHeader"))
    return t


def write_structured(doc, layout, font):
    p = layout.book.profile
    _styles(doc, font, p)
    native = p.updated(
        writing_mode="horizontal-tb",
        panels=1,
        spine=0,
        columns=1,
        rows=1,
        font_size=min(12, p.font_size),
    )
    _section(doc.sections[0], native)
    sec = doc.sections[0]
    grid = sec._sectPr.find(qn("w:docGrid"))
    if grid is not None:
        sec._sectPr.remove(grid)
    sec.top_margin = Pt(12)
    sec.bottom_margin = Pt(10)
    sec.left_margin = sec.right_margin = Pt(p.margin_x)
    for pi, page in enumerate(layout.pages):
        first = True
        for widget in page.widgets:
            if widget["kind"] == "score":
                write_score(doc, widget, font, p, layout.book.special.systems_per_page)
                continue
            if widget["kind"] == "graph":
                graph(doc, widget, font, p)
                continue
            if widget["kind"] == "table":
                table(doc, widget, font, p, layout.book.special.kind)
                continue
            para = doc.add_paragraph()
            para.paragraph_format.space_before = para.paragraph_format.space_after = Pt(
                0
            )
            para.paragraph_format.line_spacing = Pt(widget["height"])
            if first:
                para.paragraph_format.page_break_before = pi > 0
                first = False
            tag = "bamboo:meta:title" if widget["text"] == layout.book.title else None
            tagged(para, widget["text"], font, widget["size"], widget["color"], tag)
    # Column/table contents, including vertical cells, remain native text.


def import_structured(doc, saved):
    from ..structured import (
        GiftRecord,
        GiftLedger,
        FamilyPerson,
        Genealogy,
        GongcheScore,
        GongcheNote,
    )
    from ..editor import EditorSession
    from ..model import Book, Block, BambooError

    if not isinstance(saved.special, (GiftLedger, Genealogy, GongcheScore)):
        return None
    kind = saved.special.kind
    values = {}
    diagram_names = {}
    lyric_spans = {}
    order = []
    title = None
    for sdt in doc._element.body.iter(qn("w:sdt")):
        tag = sdt.find(qn("w:sdtPr"))
        tag = tag.find(qn("w:tag")) if tag is not None else None
        key = tag.get(qn("w:val"), "") if tag is not None else ""
        value = "".join(
            (n.text or "") if n.tag == qn("w:t") else "\n"
            for n in sdt.iter()
            if n.tag in {qn("w:t"), qn("w:br")}
        )
        if key == "bamboo:meta:title":
            title = title or value
        if key.startswith("bamboo:" + kind + ":"):
            _, _, record, field = key.split(":")
            if kind == "gongche" and field == "lyric":
                cell = sdt.getparent()
                while cell is not None and cell.tag != qn("w:tc"):
                    cell = cell.getparent()
                span = 1
                if cell is not None:
                    props = cell.find(qn("w:tcPr"))
                    if saved.profile.vertical:
                        merge = (
                            props.find(qn("w:vMerge")) if props is not None else None
                        )
                        if merge is not None and merge.get(qn("w:val")) == "restart":
                            row = cell.getparent()
                            nextrow = row.getnext()
                            while nextrow is not None and nextrow.tag == qn("w:tr"):
                                first = nextrow.find(qn("w:tc"))
                                vm = (
                                    first.find(qn("w:tcPr") + "/" + qn("w:vMerge"))
                                    if first is not None
                                    else None
                                )
                                if vm is None or vm.get(qn("w:val")) == "restart":
                                    break
                                span += 1
                                nextrow = nextrow.getnext()
                    else:
                        gs = props.find(qn("w:gridSpan")) if props is not None else None
                        if gs is not None:
                            span = int(gs.get(qn("w:val"), "1"))
                lyric_spans[record] = span
            if field == "diagram_name":
                diagram_names[record] = value
                continue
            if field not in {
                "name",
                "amount",
                "gift",
                "date",
                "note",
                "parents",
                "spouses",
                "birth",
                "death",
                "biography",
                "symbol",
                "lyric",
                "beat",
                "register",
                "phrase",
            }:
                continue
            if record not in values:
                values[record] = {}
                order.append(record)
            if field in values[record]:
                return None
            values[record][field] = value
    if not values and saved.special.records:
        return None
    if kind == "gongche":
        order = []
        for top in doc.tables:
            leaves = [
                t
                for t in top._tbl.iter(qn("w:tbl"))
                if len(list(t.iter(qn("w:tbl")))) == 1
            ]
            if saved.profile.vertical:
                leaves.reverse()
            for leaf in leaves:
                local = []
                phrase = None
                for tag in leaf.iter(qn("w:tag")):
                    key = tag.get(qn("w:val"), "")
                    if key.startswith("bamboo:gongche:"):
                        _, _, record, field = key.split(":")
                        if record not in local:
                            local.append(record)
                        if field == "phrase":
                            phrase = values.get(record, {}).get("phrase")
                for record in local:
                    if phrase is not None:
                        values[record]["phrase"] = phrase
                order.extend(local)
    elif saved.profile.vertical:
        order = []
        for t in doc.tables:
            local = []
            for node in t._tbl.iter(qn("w:tag")):
                key = node.get(qn("w:val"), "")
                if key.startswith("bamboo:" + kind + ":"):
                    record = key.split(":")[2]
                    if record not in local:
                        local.append(record)
            order.extend(reversed(local))
    try:
        if kind == "gift":
            records = tuple(GiftRecord(id=i, **values[i]) for i in order)
        elif kind == "gongche":
            records = tuple(
                GongcheNote(id=i, **{**values[i], "lyric_span": lyric_spans.get(i, 1)})
                for i in order
            )
        else:
            import re

            def relatives(value):
                if not value:
                    return ()
                result = []
                for label in value.split("；"):
                    match = re.match(r"^(\d+)\s+", label)
                    if not match:
                        raise BambooError("亲属编号未能识别")
                    index = int(match.group(1))
                    if not 1 <= index <= len(saved.special.records):
                        raise BambooError("亲属编号超出范围")
                    result.append(saved.special.records[index - 1].id)
                return tuple(result)

            original = {r.id: r.name for r in saved.special.records}
            for record, diagram_name in diagram_names.items():
                if record in values and diagram_name != original.get(record):
                    if values[record]["name"] not in {
                        original.get(record),
                        diagram_name,
                    }:
                        raise BambooError("人物图框与传记姓名修改冲突")
                    values[record]["name"] = diagram_name
            records = tuple(
                FamilyPerson(
                    id=i,
                    **{
                        **values[i],
                        "parents": relatives(values[i].get("parents", "")),
                        "spouses": relatives(values[i].get("spouses", "")),
                    },
                )
                for i in order
            )
        special = replace(saved.special, records=records)
        book = Book(
            title or doc.core_properties.title or saved.title,
            (Block(),),
            volume=saved.volume,
            author=doc.core_properties.author or saved.author,
            profile=saved.profile,
            special=special,
            font=saved.font,
        )
        return EditorSession(book), [
            "专用文档已按 Word 中的实际记录字段恢复，派生内容会重新计算。"
        ]
    except (BambooError, TypeError, KeyError, IndexError):
        return None


def graph(doc, widget, font, p):
    from lxml import etree
    from docx.text.paragraph import Paragraph

    anchor = doc.add_paragraph()
    anchor.paragraph_format.space_before = anchor.paragraph_format.space_after = Pt(0)
    anchor.paragraph_format.line_spacing = Pt(widget["height"])
    v = "urn:schemas-microsoft-com:vml"
    for index, line in enumerate(widget["lines"]):
        pict = element("w:pict")
        shape = etree.SubElement(
            pict,
            "{" + v + "}rect",
            id=f"JianduLink{len(doc.paragraphs)}_{index}",
            stroked="f",
            filled="t",
            fillcolor=line["color"],
        )
        left = min(line["x1"], line["x2"])
        top = min(line["y1"], line["y2"])
        width = max(line["width"], abs(line["x2"] - line["x1"]))
        height = max(line["width"], abs(line["y2"] - line["y1"]))
        shape.set(
            "style",
            f"position:absolute;margin-left:{left}pt;margin-top:{top}pt;width:{width}pt;height:{height}pt;mso-position-horizontal-relative:page;mso-position-vertical-relative:page;z-index:1",
        )
        anchor.add_run()._r.append(pict)
    for node in widget["nodes"]:
        pict = element("w:pict")
        shape = etree.SubElement(
            pict,
            "{" + v + "}rect",
            id="JianduPerson" + node.get("record", "empty"),
            strokecolor=p.rule_color,
            strokeweight=".8pt",
            fillcolor=p.paper,
        )
        shape.set(
            "style",
            f'position:absolute;margin-left:{node["x"]}pt;margin-top:{node["y"]}pt;width:{node["width"]}pt;height:{node["height"]}pt;mso-position-horizontal-relative:page;mso-position-vertical-relative:page;z-index:2',
        )
        box = etree.SubElement(
            shape,
            "{" + v + "}textbox",
            inset="4pt,4pt,4pt,4pt",
            style="layout-flow:vertical" if node["vertical"] else "",
        )
        content = element("w:txbxContent")
        box.append(content)
        pe = element("w:p")
        content.append(pe)
        para = Paragraph(pe, anchor._parent)
        para.paragraph_format.space_before = para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = Pt(node["size"] * 1.35)
        if node.get("name"):
            tagged(
                para,
                node["name"],
                font,
                node["size"],
                node["color"],
                f'bamboo:genealogy:{node["record"]}:diagram_name',
            )
            tagged(
                para,
                node["text"][len(node["name"]) :],
                font,
                node["size"],
                node["color"],
            )
        else:
            tagged(para, node["text"], font, node["size"], node["color"])
        anchor.add_run()._r.append(pict)


def write_score(doc, widget, font, p, systems):
    if widget["vertical"]:
        outer = doc.add_table(rows=1, cols=systems)
        outer.autofit = False
        for col in outer.columns:
            col.width = Pt(widget["width"] / systems)
        for cell in outer.rows[0].cells:
            cell.width = Pt(widget["width"] / systems)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = para.paragraph_format.space_after = Pt(
                0
            )
            para.paragraph_format.line_spacing = Pt(1)
            for side in ["top", "bottom"]:
                margins = cell._tc.get_or_add_tcPr().find(qn("w:tcMar"))
                if margins is None:
                    margins = element("w:tcMar")
                    cell._tc.get_or_add_tcPr().append(margins)
                margins.append(element("w:" + side, w=0, type="dxa"))
        for i, group in enumerate(widget["tables"]):
            cell = outer.cell(0, systems - i - 1)
            table(cell, group, font, p, "gongche")
            tail = cell.paragraphs[-1]
            tail.paragraph_format.line_spacing = Pt(1)
            tail.paragraph_format.space_before = tail.paragraph_format.space_after = Pt(
                0
            )
    else:
        for i, group in enumerate(widget["tables"]):
            if i:
                gap = doc.add_paragraph()
                gap.paragraph_format.line_spacing = Pt(12)
                gap.paragraph_format.space_before = gap.paragraph_format.space_after = (
                    Pt(0)
                )
            table(doc, group, font, p, "gongche")
