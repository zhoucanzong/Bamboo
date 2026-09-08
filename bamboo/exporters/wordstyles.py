"""Word adapters for reusable styles, title-slip frames and inline text seals."""

from dataclasses import replace
import math
from lxml import etree
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

from ..styles import style_registry, resolve_style


def register_styles(doc, book, font):
    from .wordflow import set_font

    for key, style in style_registry(book).items():
        name = "Bamboo S " + key
        native = (
            doc.styles[name]
            if name in doc.styles
            else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        )
        native.base_style = doc.styles["Normal"]
        set_font(
            native._element.get_or_add_rPr(),
            font.family,
            book.profile.font_size * style.font_scale,
        )
        native.font.bold = style.bold
        native.font.color.rgb = RGBColor.from_string(
            (style.ink or book.profile.ink)[1:]
        )


def format_paragraph(paragraph, style, p):
    from .wordflow import line_pitch

    f = paragraph.paragraph_format
    f.alignment = {
        "start": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "end": WD_ALIGN_PARAGRAPH.RIGHT,
    }[style.align]
    f.space_before = Pt(style.before * line_pitch(p))
    f.space_after = Pt(style.after * line_pitch(p))
    f.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    f.line_spacing = Pt(line_pitch(p))


def write_cover(doc, book, begin, end, spec, font, layout):
    from .wordflow import element, _run

    p = spec.profile
    v = "urn:schemas-microsoft-com:vml"
    width = min(spec.cover_width, p.width - 2 * p.margin_x)
    height = min(p.height * 0.68, p.body_height)
    left = (p.width - width) / 2
    top = (p.height - height) / 2
    anchor = doc.add_paragraph()
    anchor.paragraph_format.line_spacing = Pt(1)
    pict = element("w:pict")
    shape = etree.SubElement(
        pict,
        "{" + v + "}rect",
        id=f"BambooCover{begin}",
        filled="f",
        stroked="f" if spec.cover_border == "none" else "t",
    )
    color = p.border_color or resolve_style(book, book.blocks[begin]).ink or p.accent
    shape.set("strokecolor", color)
    shape.set("strokeweight", ".9pt")
    x, y = (top, p.width - left - width) if p.vertical else (left, top)
    shape.set(
        "style",
        f"position:absolute;margin-left:{x}pt;margin-top:{y}pt;width:{width}pt;height:{height}pt;"
        "mso-position-horizontal-relative:page;mso-position-vertical-relative:page;z-index:3",
    )
    if spec.cover_border == "double":
        outer_pict = element("w:pict")
        outer = etree.SubElement(
            outer_pict,
            "{" + v + "}rect",
            id=f"BambooCoverFrame{begin}",
            filled="f",
            strokecolor=color,
            strokeweight="1.4pt",
        )
        ox, oy = (
            (top - 3, p.width - left - width - 3) if p.vertical else (left - 3, top - 3)
        )
        outer.set(
            "style",
            f"position:absolute;margin-left:{ox}pt;margin-top:{oy}pt;width:{width+6}pt;height:{height+6}pt;mso-position-horizontal-relative:page;mso-position-vertical-relative:page;z-index:2",
        )
        anchor.add_run()._element.append(outer_pict)
    box = etree.SubElement(
        shape,
        "{" + v + "}textbox",
        inset="6pt,6pt,6pt,6pt",
        style="layout-flow:vertical",
    )
    content = element("w:txbxContent")
    box.append(content)
    pe = element("w:p")
    content.append(pe)
    paragraph = Paragraph(pe, anchor._parent)
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = (
        Pt(0)
    )
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(width - 12)
    for bi in range(begin, end):
        block = book.blocks[bi]
        style = resolve_style(book, block)
        glyph = next(
            (g for page in layout.pages for g in page.glyphs if g.block == bi), None
        )
        size = (
            glyph.size if glyph else min(p.font_size * style.font_scale, width * 0.65)
        )
        paragraph._element.append(
            element("w:bookmarkStart", id=bi + 100, name=f"BambooBlock{bi}")
        )
        for inline in block.inlines:
            r = _run(
                paragraph,
                inline.text,
                font,
                size,
                p.accent if inline.kind == "emphasis" else style.ink or p.ink,
                size * 1.5,
            )
            r.bold = style.bold
        paragraph._element.append(element("w:bookmarkEnd", id=bi + 100))
        if bi < end - 1:
            _run(
                paragraph,
                "　",
                font,
                max(4, size * 0.4),
                style.ink or p.ink,
                size * 0.6,
            )
    anchor.add_run()._element.append(pict)


def inline_seal(paragraph, inline, font, p, serial):
    from .wordflow import element, _run

    v = "urn:schemas-microsoft-com:vml"
    side = min(p.line_advance * 0.88, p.cell_advance * 2 * 0.88, p.font_size * 1.8)
    n = math.ceil(math.sqrt(len(inline.text)))
    inset = side * 0.1
    cell = (side - 2 * inset) / n
    pict = element("w:pict")
    shape = etree.SubElement(
        pict,
        "{" + v + "}rect",
        id=f"BambooSeal{serial}",
        filled="t" if inline.seal_style == "white" else "f",
        fillcolor=p.accent,
        strokecolor=p.accent,
        strokeweight=".8pt",
    )
    shape.set("style", f"width:{side}pt;height:{side}pt")
    shape.set("alt", "Bamboo text seal " + inline.seal_style)
    textbox = etree.SubElement(
        shape,
        "{" + v + "}textbox",
        inset=f"{inset}pt,{inset}pt,{inset}pt,{inset}pt",
        style="layout-flow:vertical",
    )
    content = element("w:txbxContent")
    textbox.append(content)
    pe = element("w:p")
    content.append(pe)
    text = Paragraph(pe, paragraph._parent)
    text.paragraph_format.space_before = text.paragraph_format.space_after = Pt(0)
    text.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    text.paragraph_format.line_spacing = Pt(cell)
    color = "#ffffff" if inline.seal_style == "white" else p.accent
    for i in range(0, len(inline.text), n):
        _run(text, inline.text[i : i + n], font, cell * 0.8, color, cell)
        if i + n < len(inline.text):
            text.add_run().add_break()
    paragraph.add_run()._element.append(pict)
