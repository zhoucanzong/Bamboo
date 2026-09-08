"""Direction-preserving native Word flow and optional facsimile export."""

from io import BytesIO
import json
import hashlib
import uuid
from dataclasses import replace

import fitz
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

from .pdf import pdf_document
from ..model import BambooError, Book, Block
from ..layout import _frame


def _embed_font(doc, font):
    if font.rights & 4 and not font.rights & 8:
        raise BambooError("字体仅允许预览与打印嵌入，不能生成带嵌入字体的可编辑稿")
    key = uuid.UUID(bytes=hashlib.sha256(font.data).digest()[:16])
    mask = key.bytes[::-1]
    data = bytearray(font.data)
    for i in range(min(32, len(data))):
        data[i] ^= mask[i % 16]
    table = doc.part.part_related_by(RT.FONT_TABLE)
    # python-docx represents the font table as a generic OPC part.
    root = etree.fromstring(table.blob)
    entry = next((n for n in root if n.get(qn("w:name")) == font.family), None)
    if entry is None:
        entry = OxmlElement("w:font")
        entry.set(qn("w:name"), font.family)
        root.append(entry)
    font_part = Part(
        PackURI("/word/fonts/bamboo.odttf"),
        "application/vnd.openxmlformats-officedocument.obfuscatedFont",
        bytes(data),
        doc.part.package,
    )
    rel = table.relate_to(font_part, RT.FONT)
    embedded = OxmlElement("w:embedRegular")
    embedded.set(qn("r:id"), rel)
    embedded.set(qn("w:fontKey"), "{" + str(key).upper() + "}")
    embedded.set(qn("w:subsetted"), "true" if not font.rights & 0x100 else "false")
    entry.append(embedded)
    table._blob = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )


def _source_part(doc, layout):
    root = etree.Element("{urn:bamboo:document:1}source")
    root.text = json.dumps(layout.to_dict()["book"], ensure_ascii=False)
    part = Part(
        PackURI("/customXml/bamboo-source.xml"),
        "application/xml",
        etree.tostring(root, xml_declaration=True, encoding="UTF-8"),
        doc.part.package,
    )
    doc.part.relate_to(part, RT.CUSTOM_XML)


def _floating_picture(
    paragraph, image, width, height, description, x=0, y=0, wrap=False
):
    shape = paragraph.add_run().add_picture(image, width=Pt(width), height=Pt(height))
    inline = shape._inline
    anchor = OxmlElement("wp:anchor")
    for key, value in {
        "distT": "0",
        "distB": "0",
        "distL": "0",
        "distR": "0",
        "simplePos": "0",
        "relativeHeight": "0",
        "behindDoc": "1",
        "locked": "0",
        "layoutInCell": "1",
        "allowOverlap": "1",
    }.items():
        anchor.set(key, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    for name, coordinate in (("wp:positionH", x), ("wp:positionV", y)):
        pos = OxmlElement(name)
        pos.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset")
        offset.text = str(round(coordinate * 12700))
        pos.append(offset)
        anchor.append(pos)
    anchor.append(inline.find(qn("wp:extent")))
    wrapping = OxmlElement("wp:wrapSquare" if wrap else "wp:wrapNone")
    if wrap:
        wrapping.set("wrapText", "bothSides")
        anchor.set("behindDoc", "0")
    anchor.append(wrapping)
    props = inline.find(qn("wp:docPr"))
    props.set("descr", description)
    anchor.append(props)
    frame = inline.find(qn("wp:cNvGraphicFramePr"))
    if frame is not None:
        anchor.append(frame)
    anchor.append(inline.find(qn("a:graphic")))
    inline.getparent().replace(inline, anchor)


def _flow_decoration(doc, layout, font, section=None):
    """Only stationery is an image. All body and annotation stories stay text.

    A wrapping spine object in the repeating header excludes the central band
    from the continuous vertical story, without forced page/column breaks.
    """
    p = layout.book.profile
    glyphs, lines, polygons = _frame(layout.book, 1)
    if (
        not glyphs
        and not lines
        and not polygons
        and p.paper == "#ffffff"
        and not p.spine
    ):
        return
    glyphs = [g for g in glyphs if g.y < p.margin_top + p.body_height * 0.83]
    leaf = replace(
        layout.pages[0],
        glyphs=tuple(glyphs),
        lines=tuple(lines),
        polygons=tuple(polygons),
    )
    decoration = replace(layout, pages=(leaf,))
    header = (section or doc.sections[0]).header
    para = header.paragraphs[0]
    para.paragraph_format.space_before = para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    para.paragraph_format.line_spacing = Pt(1)
    with pdf_document(decoration, font) as pdf:
        page = pdf[0]
        _floating_picture(
            para,
            BytesIO(page.get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes("png")),
            p.width,
            p.height,
            "版框与版心装饰",
        )
        if p.spine and p.panels == 2 and p.vertical:
            x = p.margin_x + p.panel_width
            # Cover the full inline extent: otherwise a reader may flow a last
            # character through a tiny gap below the spine's body rectangle.
            rect = fitz.Rect(x, 0, x + p.spine, p.height)
            sprite = page.get_pixmap(matrix=fitz.Matrix(2, 2), clip=rect)
            _floating_picture(
                para,
                BytesIO(sprite.tobytes("png")),
                p.spine,
                p.height,
                "版心避让区",
                x=x,
                y=0,
                wrap=True,
            )
    # A native PAGE field, never a baked-in page number.
    v = "urn:schemas-microsoft-com:vml"
    if p.spine and p.show_page_number:
        from .wordflow import set_font

        style = (
            doc.styles["Bamboo Page Number"]
            if "Bamboo Page Number" in doc.styles
            else doc.styles.add_style("Bamboo Page Number", 1)
        )
        set_font(style._element.get_or_add_rPr(), font.family, 11)
        style.paragraph_format.line_spacing = Pt(14)
        x = p.margin_x + (p.panel_width if p.panels == 2 else 0)
        y = p.margin_top + p.body_height * 0.83
        pict = OxmlElement("w:pict")
        shape = etree.SubElement(
            pict,
            "{" + v + "}rect",
            id="BambooPageNumber" + str(len(doc.sections)),
            stroked="f",
            filled="f",
        )
        shape.set(
            "style",
            f"position:absolute;margin-left:{x}pt;margin-top:{y}pt;width:{p.spine}pt;height:40pt;"
            "mso-position-horizontal-relative:page;mso-position-vertical-relative:page;z-index:2",
        )
        box = etree.SubElement(shape, "{" + v + "}textbox", inset="0,0,0,0")
        content = OxmlElement("w:txbxContent")
        box.append(content)
        pp = OxmlElement("w:p")
        content.append(pp)
        prop = OxmlElement("w:pPr")
        style_ref = OxmlElement("w:pStyle")
        style_ref.set(qn("w:val"), style.style_id)
        prop.append(style_ref)
        align = OxmlElement("w:jc")
        align.set(qn("w:val"), "center")
        prop.append(align)
        pp.append(prop)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), " PAGE ")
        rr = OxmlElement("w:r")
        rp = OxmlElement("w:rPr")
        from .wordflow import set_font

        set_font(rp, font.family, 11)
        rr.append(rp)
        t = OxmlElement("w:t")
        t.text = "一"
        rr.append(t)
        field.append(rr)
        pp.append(field)
        para.add_run()._element.append(pict)


def export_docx(layout, font, path, mode="flow", dpi=180):
    doc = Document()
    p = layout.book.profile
    section = doc.sections[0]
    section.page_width, section.page_height = Pt(p.width), Pt(p.height)
    doc.core_properties.title = layout.book.title
    doc.core_properties.author = layout.book.author
    doc.core_properties.subject = layout.book.volume
    doc.core_properties.comments = "Created with Bamboo"
    if mode == "facsimile":
        section.top_margin = section.bottom_margin = section.left_margin = (
            section.right_margin
        ) = Pt(0)
        section.header_distance = section.footer_distance = Pt(0)
        with pdf_document(layout, font) as pdf:
            for i, page in enumerate(pdf):
                from docx.enum.section import WD_SECTION_START

                changed_size = i > 0 and (
                    abs(section.page_width.pt - page.rect.width) > 0.1
                    or abs(section.page_height.pt - page.rect.height) > 0.1
                )
                if changed_size:
                    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
                section.page_width, section.page_height = Pt(page.rect.width), Pt(
                    page.rect.height
                )
                para = doc.add_paragraph()
                para.paragraph_format.page_break_before = i > 0 and not changed_size
                para.paragraph_format.space_before = (
                    para.paragraph_format.space_after
                ) = Pt(0)
                para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                para.paragraph_format.line_spacing = Pt(1)
                image = BytesIO(
                    page.get_pixmap(
                        matrix=fitz.Matrix(dpi / 72, dpi / 72), alpha=False
                    ).tobytes("png")
                )
                _floating_picture(
                    para,
                    image,
                    page.rect.width,
                    page.rect.height,
                    f"{layout.book.title} 第 {i+1} 葉，古籍保真版",
                )
    elif mode in {"flow", "editable"}:
        from .wordflow import write_flow

        def decorate(section, begin, end, spec):
            info = next(s for s in layout.sections if s["block_start"] == begin)
            page = layout.pages[info["page_offset"]]
            frame_book = Book(
                spec.title if spec.title is not None else layout.book.title,
                (Block(),),
                volume=spec.volume or "",
                author=spec.author or "",
                profile=spec.profile,
            )
            frame_layout = replace(
                layout, book=frame_book, pages=(replace(page, profile=spec.profile),)
            )
            _flow_decoration(doc, frame_layout, font, section)

        write_flow(doc, layout, font, decorate)
        _embed_font(doc, font)
    else:
        raise ValueError("DOCX mode must be flow or facsimile")
    _source_part(doc, layout)
    doc.save(str(path))
